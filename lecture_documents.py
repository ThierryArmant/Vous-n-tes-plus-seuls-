import csv
import io
import zipfile
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 60000
MAX_ROWS = 500
MAX_COLS = 40


def lire_document(name, data):
    if len(data) > MAX_BYTES:
        raise ValueError('Fichier trop volumineux : maximum 10 Mo par fichier.')
    ext = Path(name).suffix.lower()
    parts = []
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('Ce PDF est protégé par mot de passe.')
        if len(reader.pages) > 100:
            raise ValueError('Ce PDF dépasse 100 pages. Découpez-le avant de le transmettre.')
        for index, page in enumerate(reader.pages, 1):
            texte = (page.extract_text(extraction_mode='layout') or '').strip()
            if not texte:
                raise ValueError(f'La page {index} ne contient pas de texte lisible. Les PDF numérisés nécessitent une reconnaissance OCR, pas encore disponible.')
            parts.append(f'Page {index} :\n{texte}')
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Classeur trop volumineux une fois décompressé.')
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        total = 0
        try:
            for sheet in workbook.worksheets:
                if sheet.max_column and sheet.max_column > MAX_COLS:
                    raise ValueError('Le classeur dépasse 40 colonnes par feuille.')
                parts.append(f'Feuille : {sheet.title}')
                for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    total += 1
                    if total > MAX_ROWS:
                        raise ValueError('Le classeur dépasse 500 lignes. Découpez-le pour une analyse complète.')
                    parts.append(f'Ligne {index} : ' + ' | '.join('' if value is None else str(value) for value in row))
        finally:
            workbook.close()
    elif ext == '.csv':
        texte = data.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(texte[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        for index, row in enumerate(csv.reader(io.StringIO(texte), dialect), 1):
            if index > MAX_ROWS or len(row) > MAX_COLS:
                raise ValueError('Le CSV dépasse 500 lignes ou 40 colonnes. Découpez-le pour une analyse complète.')
            parts.append(f'Ligne {index} : ' + ' | '.join(row))
    else:
        raise ValueError('Format non pris en charge.')
    text = '\n'.join(parts)
    if not text.strip():
        raise ValueError('Ce fichier ne contient pas de texte lisible.')
    if len(text) > MAX_TEXT:
        raise ValueError('Le texte extrait dépasse 60 000 caractères. Découpez le fichier pour une analyse complète.')
    return text
